"""
Модуль «AI Офферы»: загрузка таблиц лидов, генерация офферов нейросетью,
учёт результатов звонков. Доступен только администраторам.
"""
import csv
import io
import queue
import re
import threading
import time
from datetime import datetime
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app import config
from app.ai_client import AIGenerationError, generate_pitch, is_configured
from app.database import SessionLocal, get_db
from app.models import CallLead
from app.routers.auth import require_admin
from app.schemas import CallLeadOut, CallLeadUpdate, OffersStats

router = APIRouter(prefix="/api/offers", tags=["AI Офферы"], dependencies=[Depends(require_admin)])

ALLOWED_CALL_STATUSES = ["Новый", "Согласие", "Отказ", "Перезвонить", "Не дозвонился"]

# Ключевые слова для автоопределения колонок таблицы
COLUMN_PATTERNS = {
    "company_name": ["назван", "компан", "организац", "фирм", "наименован", "имя", "company", "name", "organization", "firm"],
    "phone": ["телефон", "тел", "номер", "мобиль", "связ", "phone", "tel", "mobile", "contact"],
    "city": ["город", "населён", "населен", "регион", "city"],
    "niche": ["ниш", "категор", "сфер", "вид деят", "рубрик", "специализ", "niche", "category"],
    "address": ["адрес", "address", "ул.", "улица"],
    "website": ["сайт", "домен", "web", "url", "ресурс"],
}

# Заголовки колонок-«анализа сайта»: там лежат статусы, а не ссылки
# (например, «Проверка сайта» → «🚫 Нет сайта в карточке»).
WEBSITE_HEADER_STOP_WORDS = (
    "проверк", "анализ", "слабые", "статус", "оценк", "есть сайт",
    "наличие сайт", "нет сайт", "без сайт", "отсутств",
)

# Слова-маркеры статусов: значение с ними сайтом считать нельзя.
WEBSITE_VALUE_STOP_RE = re.compile(
    r"(нет\s*сайт|без\s*сайт|отсутств|не\s*указан|слабые\s*места|проверк|конструктор|"
    r"поддомен|не\s*работает|брошен|ошибк|на\s*продаже|заглушк)",
    re.IGNORECASE,
)

# Домен: латиница и кириллица (зоны .рф, .рус), отдельные метки через точку.
DOMAIN_CANDIDATE_RE = re.compile(
    r"(?<![A-Za-z0-9\u0400-\u04FF-])"
    r"((?:[A-Za-z0-9\u0400-\u04FF](?:[A-Za-z0-9\u0400-\u04FF-]*[A-Za-z0-9\u0400-\u04FF])?\.)+"
    r"[A-Za-z\u0400-\u04FF][A-Za-z0-9\u0400-\u04FF-]{1,})"
    r"(?![A-Za-z0-9\u0400-\u04FF-])"
)

# Домены без схемы (http://…) признаём сайтом только в популярных зонах.
TRUSTED_TLDS = {
    "ru", "рф", "рус", "москва", "com", "net", "org", "su", "io", "biz", "info",
    "site", "online", "shop", "store", "pro", "club", "tech", "space", "life",
    "top", "xyz", "me", "app", "dev", "cloud", "media", "studio", "agency",
    "digital", "market", "expert", "company", "business", "name", "mobi", "tv",
    "kz", "by", "ua", "uz", "am", "ge", "md", "az", "kg", "tj",
    "com.ru", "ru.com", "com.ua", "xn--p1ai", "xn--80asehdb",
}

# Агрегаторы, карты и мессенджеры: их ссылки сайтом компании не считаем.
BLOCKED_DOMAINS = {
    "2gis.ru", "2gis.com", "2gis.kz", "2gis.by", "yandex.ru", "yandex.com", "ya.ru",
    "google.com", "google.ru", "goo.gl", "maps.google.com", "g.page", "avito.ru",
    "wa.me", "api.whatsapp.com", "whatsapp.com", "t.me", "telegram.me", "telegram.org",
    "vk.com", "vk.ru", "instagram.com", "facebook.com", "fb.me", "ok.ru",
    "youtube.com", "youtu.be", "dzen.ru", "ozon.ru", "wildberries.ru",
    "market.yandex.ru", "maps.yandex.ru",
}


def _normalize_header(value) -> str:
    return str(value or "").strip().lower()


def _match_column(header: str, kind: str) -> bool:
    # Колонки-«анализ сайта» («Проверка сайта», «Слабые места», «Есть сайт»)
    # содержат текстовые статусы, а не ссылку — сайтом их считать нельзя.
    if kind == "website" and any(stop in header for stop in WEBSITE_HEADER_STOP_WORDS):
        return False
    return any(pattern in header for pattern in COLUMN_PATTERNS[kind])


def _cell_str(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in ("none", "nan", "null") else text


def _cell_phone(value) -> str:
    """Вытаскивает телефон из значения (число, «+7 999 …», несколько номеров)."""
    text = _cell_str(value)
    if not text:
        return ""
    # Оставляем цифры, +, запятые/точки с запятой как разделители нескольких номеров
    cleaned = "".join(ch for ch in text if ch.isdigit() or ch in "+,;")
    return cleaned[:200] or text[:100]


def _extract_domain(text: str) -> str:
    """Возвращает первый реальный домен из текста или пустую строку."""
    for match in DOMAIN_CANDIDATE_RE.finditer(text or ""):
        domain = match.group(1).lower()
        if domain.startswith("www."):
            domain = domain[4:]
        tld = domain.rsplit(".", 1)[-1]
        if domain in BLOCKED_DOMAINS or any(domain.endswith("." + blocked) for blocked in BLOCKED_DOMAINS):
            continue
        prefix = (text or "")[max(0, match.start() - 8):match.start()].lower()
        has_scheme = "http" in prefix or prefix.endswith("//") or prefix.endswith("www.")
        if has_scheme or tld in TRUSTED_TLDS:
            return domain
    return ""


def _cell_website(value, strict: bool = True) -> str:
    """
    Достаёт из ячейки сайт компании.

    В выгрузках рядом с сайтом часто идут колонки-статусы («Проверка сайта»,
    «Слабые места»): их значения вида «🚫 Нет сайта в карточке» или
    «🛠️ Слабые места · 40/100» сайтом не являются, иначе панель показывает
    «есть сайт» там, где его нет.
    """
    text = _cell_str(value)
    if not text:
        return ""
    has_scheme = bool(re.search(r"https?://|www\.", text, re.IGNORECASE))
    if strict and not has_scheme and WEBSITE_VALUE_STOP_RE.search(text):
        return ""
    return _extract_domain(text)


def _find_website_in_row(row: List[str], mapping: dict) -> str:
    """
    Ищет домен компании в остальных колонках строки.

    Нужно, когда колонка «Сайт» в выгрузке — это статус, а настоящий адрес
    сайта упомянут в другой колонке (например, в готовом оффере или контактах).
    """
    skip = set(mapping.values())
    for idx, cell in enumerate(row):
        if idx in skip:
            continue
        domain = _cell_website(cell, strict=False)
        if domain:
            return domain
    return ""


def parse_table_rows(content: bytes, filename: str) -> List[dict]:
    """Разбирает xlsx/xls/csv в список dict с нормализованными полями."""
    lower_name = filename.lower()

    rows: List[List[str]] = []
    header_row: List[str] = []

    if lower_name.endswith((".xlsx", ".xls")):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise HTTPException(status_code=500, detail="На сервере не установлена библиотека openpyxl")

        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(status_code=400, detail="Не удалось открыть Excel-файл. Сохраните его в формате .xlsx")

        sheet = workbook.active
        for row in sheet.iter_rows(values_only=True):
            rows.append(["" if v is None else str(v) for v in row])
        workbook.close()
    else:
        # CSV: пробуем utf-8, затем cp1251 (частый случай для выгрузок из 1С/Excel)
        text = None
        for encoding in ("utf-8-sig", "utf-8", "cp1251"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise HTTPException(status_code=400, detail="Не удалось определить кодировку CSV-файла")

        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows = [row for row in reader]

    rows = [row for row in rows if any(_cell_str(c) for c in row)]
    if not rows:
        raise HTTPException(status_code=400, detail="Файл пустой или не содержит данных")

    # Определяем строку заголовка: если в первой строке есть телефонный/текстовый заголовок
    first_normalized = [_normalize_header(c) for c in rows[0]]
    looks_like_header = any(_match_column(h, kind) for h in first_normalized for kind in ("phone", "company_name", "website"))

    mapping = {}
    if looks_like_header:
        header_row = rows[0]
        data_rows = rows[1:]
        used = set()
        # Сначала специфичные колонки (телефон/сайт/город/ниша/адрес), затем название
        for kind in ("phone", "website", "city", "niche", "address", "company_name"):
            for idx, h in enumerate(first_normalized):
                if idx in used or not h:
                    continue
                if _match_column(h, kind):
                    mapping[kind] = idx
                    used.add(idx)
                    break
        # Название: если не нашли по заголовку — первая неиспользованная текстовая колонка
        if "company_name" not in mapping:
            for idx in range(len(header_row)):
                if idx not in used:
                    mapping["company_name"] = idx
                    used.add(idx)
                    break
    else:
        # Заголовка нет: первая заполненная колонка — название, любая колонка с телефоноподобным содержимым — телефон
        data_rows = rows
        mapping = {"company_name": 0}
        phone_idx = None
        for idx in range(1, len(rows[0])):
            sample_values = [r[idx] for r in rows[1:6] if len(r) > idx]
            digits = sum(1 for v in sample_values if sum(ch.isdigit() for ch in v) >= 6)
            if sample_values and digits >= max(1, len(sample_values) // 2):
                phone_idx = idx
                break
        if phone_idx is not None:
            mapping["phone"] = phone_idx

    leads: List[dict] = []
    for row in data_rows:
        def get(kind: str) -> str:
            idx = mapping.get(kind)
            if idx is None or idx >= len(row):
                return ""
            return _cell_str(row[idx])

        company = get("company_name")
        phone = _cell_phone(row[mapping["phone"]]) if "phone" in mapping and mapping["phone"] < len(row) else ""
        if not company and not phone:
            continue  # Полностью пустая смысла строка

        # Сайт: берём только реальный домен; если в колонке «Сайт» статус —
        # ищем адрес сайта в остальных колонках строки.
        website = _cell_website(get("website"))
        if not website:
            website = _find_website_in_row(row, mapping)

        # Прочие неиспользованные колонки — в подсказку для нейросети
        used_indexes = set(mapping.values())
        extra_parts = []
        if header_row:
            for idx, cell in enumerate(row):
                if idx in used_indexes or idx >= len(header_row):
                    continue
                value = _cell_str(cell)
                if value and _normalize_header(header_row[idx]):
                    extra_parts.append(f"{header_row[idx].strip()}: {value}")
        else:
            for idx, cell in enumerate(row):
                if idx in used_indexes:
                    continue
                value = _cell_str(cell)
                if value:
                    extra_parts.append(value)

        leads.append({
            "company_name": company or (phone or "Без названия"),
            "phone": phone,
            "city": get("city"),
            "niche": get("niche"),
            "address": get("address"),
            "website": website,
            "extra_info": "; ".join(extra_parts)[:1000] or None,
        })

    if not leads:
        raise HTTPException(status_code=400, detail="В таблице не найдено строк с названиями компаний или телефонами")
    if len(leads) > config.UPLOAD_MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"Слишком большая таблица: максимум {config.UPLOAD_MAX_ROWS} строк за загрузку")
    return leads


@router.post("/upload", response_model=dict)
async def upload_table(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Загружает Excel/CSV таблицу лидов и сохраняет строки в базу."""
    if not file.filename:
        raise HTTPException(status_code=400, detail="Файл не выбран")

    lower_name = file.filename.lower()
    if not lower_name.endswith((".xlsx", ".xls", ".csv")):
        raise HTTPException(status_code=400, detail="Поддерживаются форматы .xlsx, .xls и .csv")

    content = await file.read()
    if len(content) > config.UPLOAD_MAX_SIZE_MB * 1024 * 1024:
        raise HTTPException(status_code=400, detail=f"Файл больше {config.UPLOAD_MAX_SIZE_MB} МБ")

    leads = parse_table_rows(content, file.filename)

    objects = [
        CallLead(source_file=file.filename, **lead)
        for lead in leads
    ]
    db.add_all(objects)
    db.commit()

    return {
        "status": "ok",
        "inserted": len(objects),
        "source_file": file.filename,
        "message": f"Загружено {len(objects)} компаний из «{file.filename}»"
    }


@router.get("/leads", response_model=List[CallLeadOut])
def get_leads(
    status: Optional[str] = Query(None, description="Фильтр по статусу звонка"),
    search: Optional[str] = Query(None, description="Поиск по названию/телефону/нише"),
    only_without_offer: bool = Query(False, description="Только лиды без оффера"),
    limit: int = Query(500, ge=1, le=2000),
    db: Session = Depends(get_db)
):
    query = db.query(CallLead)
    if status and status != "Все":
        query = query.filter(CallLead.call_status == status)
    if only_without_offer:
        query = query.filter(CallLead.offer_text.is_(None))
    if search:
        s = f"%{search.strip()}%"
        query = query.filter(or_(
            CallLead.company_name.ilike(s),
            CallLead.phone.ilike(s),
            CallLead.niche.ilike(s),
            CallLead.city.ilike(s),
        ))
    return query.order_by(CallLead.id.asc()).offset(0).limit(limit).all()


@router.get("/stats", response_model=OffersStats)
def get_stats(db: Session = Depends(get_db)):
    total = db.query(CallLead).count()
    with_offer = db.query(CallLead).filter(CallLead.offer_text.isnot(None)).count()

    status_counts = dict(db.query(CallLead.call_status, func.count(CallLead.id)).group_by(CallLead.call_status).all())

    return OffersStats(
        total=total,
        with_offer=with_offer,
        without_offer=total - with_offer,
        by_status=status_counts
    )


# ---------- Фоновая генерация офферов (строго по одному) ----------
#
# Нейросети нужно 1–3 минуты на один оффер, поэтому запрос не ждёт ответ:
# POST /generate мгновенно ставит лид в очередь, а отдельный поток-обработчик
# генерирует офферы строго последовательно. Так массовая генерация не упирается
# в таймауты хостинга и не заваливает шлюз параллельными запросами.

_generation_queue: "queue.Queue[int]" = queue.Queue()
_generation_lock = threading.Lock()
_generation_state = {
    "pending": [],    # id лидов, ждущих очереди
    "current": None,  # id лида, который генерируется прямо сейчас
    "done": 0,        # готово с момента запуска задачи
    "failed": 0,      # ошибок с момента запуска задачи
}
_worker_started = False


def _generation_progress() -> dict:
    with _generation_lock:
        return {
            "pending": list(_generation_state["pending"]),
            "current": _generation_state["current"],
            "done": _generation_state["done"],
            "failed": _generation_state["failed"],
        }


def _generate_pitch_with_retry(lead_data: dict, attempts: int = 3) -> dict:
    """
    Вызывает нейросеть с повторами для временных сбоев
    (таймауты, 429/5xx от шлюза). Между попытками — пауза.
    """
    last_error: Optional[AIGenerationError] = None
    for attempt in range(1, attempts + 1):
        try:
            return generate_pitch(lead_data)
        except AIGenerationError as exc:
            last_error = exc
            message = str(exc)
            transient = "Сеть недоступна" in message or re.search(r"API вернул (429|5\d\d)", message) is not None
            if not transient or attempt == attempts:
                raise
            time.sleep(10 * attempt)
    raise last_error or AIGenerationError("Неизвестная ошибка генерации")


def _process_lead_offer(db: Session, lead_id: int) -> bool:
    """Генерирует и сохраняет оффер одного лида. Возвращает успех."""
    lead = db.query(CallLead).filter(CallLead.id == lead_id).first()
    if not lead:
        return False

    lead_data = {
        "company_name": lead.company_name,
        "phone": lead.phone,
        "city": lead.city,
        "niche": lead.niche,
        "address": lead.address,
        "website": lead.website,
        "extra_info": lead.extra_info,
    }

    try:
        pitch = _generate_pitch_with_retry(lead_data)
    except AIGenerationError as exc:
        lead.offer_error = str(exc)[:500]
        db.commit()
        return False

    lead.offer_hook = pitch["hook"]
    lead.offer_text = pitch["offer"]
    lead.offer_objections = pitch["objections"]
    lead.offer_closing = pitch["closing"]
    lead.offer_error = None
    lead.offer_generated_at = datetime.utcnow()
    db.commit()
    return True


def _generation_worker() -> None:
    """Фоновый поток: забирает лиды из очереди и обрабатывает по одному."""
    while True:
        lead_id = _generation_queue.get()
        with _generation_lock:
            if lead_id in _generation_state["pending"]:
                _generation_state["pending"].remove(lead_id)
            _generation_state["current"] = lead_id

        db = SessionLocal()
        try:
            success = _process_lead_offer(db, lead_id)
        except Exception:  # noqa: BLE001 — поток не должен умирать
            success = False
            try:
                db.rollback()
                stale = db.query(CallLead).filter(CallLead.id == lead_id).first()
                if stale:
                    stale.offer_error = "Внутренняя ошибка при генерации, попробуйте ещё раз"
                    db.commit()
            except Exception:  # noqa: BLE001
                pass
        finally:
            db.close()
            with _generation_lock:
                _generation_state["current"] = None
                if success:
                    _generation_state["done"] += 1
                else:
                    _generation_state["failed"] += 1
            _generation_queue.task_done()


def _ensure_worker() -> None:
    global _worker_started
    with _generation_lock:
        if _worker_started:
            return
        _worker_started = True
    threading.Thread(target=_generation_worker, daemon=True).start()


@router.get("/status")
def generation_status():
    """Прогресс фоновой генерации: сколько в очереди, готово, ошибок."""
    return _generation_progress()


@router.post("/leads/{lead_id}/generate")
def generate_offer(lead_id: int, db: Session = Depends(get_db)):
    """
    Ставит лид в очередь на генерацию оффера и сразу отвечает.
    Сама генерация идёт в фоне, строго по одному лиду: клиент следит
    за прогрессом через GET /api/offers/status.
    """
    if not is_configured():
        raise HTTPException(status_code=503, detail="Нейросеть не настроена на сервере (AI_API_URL/AI_API_KEY)")

    lead = db.query(CallLead).filter(CallLead.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Лид не найден")

    progress = _generation_progress()
    if lead_id == progress["current"] or lead_id in progress["pending"]:
        return {"status": "queued", "message": "Лид уже в очереди на генерацию", "progress": progress}

    lead.offer_error = None
    db.commit()

    with _generation_lock:
        _generation_state["pending"].append(lead_id)
    _generation_queue.put(lead_id)
    _ensure_worker()

    return {
        "status": "queued",
        "message": "Оффер поставлен в очередь, генерация идёт в фоне",
        "progress": _generation_progress(),
    }


@router.patch("/leads/{lead_id}", response_model=CallLeadOut)
def update_lead(lead_id: int, update: CallLeadUpdate, db: Session = Depends(get_db)):
    """Проставляет итог звонка (согласие/отказ/перезвонить) и комментарий."""
    lead = db.query(CallLead).filter(CallLead.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Лид не найден")

    if update.call_status is not None:
        if update.call_status not in ALLOWED_CALL_STATUSES:
            raise HTTPException(status_code=400, detail=f"Недопустимый статус. Допустимо: {', '.join(ALLOWED_CALL_STATUSES)}")
        lead.call_status = update.call_status
        lead.called_at = None if update.call_status == "Новый" else datetime.utcnow()

    if update.call_notes is not None:
        lead.call_notes = update.call_notes.strip() or None

    db.commit()
    db.refresh(lead)
    return lead


@router.delete("/leads/{lead_id}")
def delete_lead(lead_id: int, db: Session = Depends(get_db)):
    lead = db.query(CallLead).filter(CallLead.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Лид не найден")
    db.delete(lead)
    db.commit()
    return {"status": "ok", "message": f"Лид #{lead_id} удалён"}


@router.post("/clear")
def clear_all(db: Session = Depends(get_db)):
    """Полная очистка базы лидов обзвона."""
    deleted = db.query(CallLead).delete()
    db.commit()
    return {"status": "ok", "deleted": deleted, "message": f"База обзвона очищена (удалено {deleted})"}


def repair_lead_websites(db: Session) -> int:
    """
    Чистит уже сохранённые сайты лидов от текстовых статусов.

    Раньше в поле «Сайт» попадали значения колонок-анализов
    («🚫 Нет сайта в карточке», «🛠️ Слабые места · 40/100»), из-за чего
    панель показывала «есть сайт» у компаний без сайта. Функция удаляет
    такие значения и, если домен упомянут в доп. информации, восстанавливает
    его. Идемпотентна: повторный запуск ничего не меняет.

    Возвращает количество исправленных лидов.
    """
    changed = 0
    for lead in db.query(CallLead).all():
        current = (lead.website or "").strip()
        cleaned = _cell_website(lead.website) or _cell_website(lead.extra_info, strict=False)
        if cleaned == current:
            continue
        lead.website = cleaned or None
        changed += 1
    if changed:
        db.commit()
    return changed
